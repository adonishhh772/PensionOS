from __future__ import annotations

import io
from typing import Iterator

from openpyxl import load_workbook

from data_onboarding.adapters.common import filter_delta, infer_schema, read_control_totals
from data_onboarding.adapters.ports import AdapterPort


class XLSXAdapter(AdapterPort):
    def __init__(self, sheet_index: int = 0) -> None:
        self.sheet_index = sheet_index

    def inspect(self, source_meta: dict) -> dict:
        return {"format": "xlsx", "sheet_index": self.sheet_index}

    def read_schema(self, source_stream: bytes) -> dict:
        return infer_schema(self.stream_records(source_stream))

    def stream_records(self, source_stream: bytes) -> Iterator[dict]:
        bio = io.BytesIO(source_stream)
        wb = load_workbook(bio, read_only=True)
        try:
            ws = wb[wb.sheetnames[self.sheet_index]]
            rows = ws.iter_rows(values_only=True)
            header = None
            for row in rows:
                if header is None:
                    header = [str(cell) if cell is not None else "" for cell in row]
                    continue
                yield {header[i]: row[i] for i in range(len(header)) if i < len(row)}
        finally:
            wb.close()

    def read_control_totals(self, source_stream: bytes) -> dict:
        return read_control_totals(self.stream_records(source_stream))

    def read_delta(self, source_stream: bytes, since_checkpoint: str | None = None) -> Iterator[dict]:
        return filter_delta(self.stream_records(source_stream), since_checkpoint)
